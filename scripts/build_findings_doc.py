"""Build the findings handout from saved results with the bundled Python runtime."""

import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data/processed/chapter7"


def read(name):
    return json.loads((DATA / name).read_text(encoding="utf-8"))


def main():
    report = read("chunk_report.json")
    main_run = read("benchmark_jina_dev_child_pro_all_source_audited_v2.json")
    neighbors = read("benchmark_jina_dev_neighbors_pro_hybrid_source_audited_v2.json")
    parents = read("benchmark_jina_dev_parent_pro_hybrid_source_audited_v2.json")
    usage = read("benchmark_jina_dev_child_pro.json")["api_usage_this_run"]
    manifest = read("manifest.json")
    doc = Document()
    doc.core_properties.title = "GST RAG Benchmark Findings"
    doc.core_properties.subject = "Chunking and parent context retrieval experiment"
    doc.core_properties.author = "GST RAG Project"
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(.7)
    section.left_margin = section.right_margin = Inches(.75)
    section.footer_distance = Inches(.3)
    for name in ["Normal", "Title", "Subtitle", "Heading 1", "Heading 2"]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_after = Pt(6)
    doc.styles["Normal"].font.size = Pt(11)
    doc.styles["Normal"].paragraph_format.line_spacing = 1.08
    doc.styles["Title"].font.size = Pt(25)
    doc.styles["Title"].font.bold = True
    doc.styles["Heading 1"].font.size = Pt(14)
    doc.styles["Heading 1"].paragraph_format.space_before = Pt(12)
    doc.styles["Heading 2"].font.size = Pt(11.5)
    doc.styles["Heading 2"].paragraph_format.space_before = Pt(8)
    footer = section.footer.paragraphs[0]
    footer.paragraph_format.tab_stops.add_tab_stop(Inches(6.65))
    footer.add_run("GST RAG experiment findings\t").font.size = Pt(9)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)

    def p(text, small=False):
        paragraph = doc.add_paragraph(text)
        if small:
            for run in paragraph.runs:
                run.font.size = Pt(9.5)
        return paragraph

    def heading(text):
        doc.add_heading(text, level=1)

    def table(headers, rows, widths):
        tab = doc.add_table(rows=1, cols=len(headers))
        tab.alignment = WD_TABLE_ALIGNMENT.CENTER
        tab.autofit = False
        for column, width in zip(tab.columns, widths):
            column.width = Inches(width)
        for cell, name in zip(tab.rows[0].cells, headers):
            cell.text = name
        for row in rows:
            for cell, value in zip(tab.add_row().cells, row):
                cell.text = str(value)
        repeat = OxmlElement("w:tblHeader")
        tab.rows[0]._tr.get_or_add_trPr().append(repeat)
        for i, row in enumerate(tab.rows):
            # Keep each row intact; permit pagination between rows if needed.
            cant_split = OxmlElement("w:cantSplit")
            row._tr.get_or_add_trPr().append(cant_split)
            for j, cell in enumerate(row.cells):
                cell.width = Inches(widths[j])
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                tcpr = cell._tc.get_or_add_tcPr()
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "DDE8F0" if i == 0 else "F5F6F7" if i % 2 == 0 else "FFFFFF")
                tcpr.append(shading)
                borders = OxmlElement("w:tcBorders")
                for edge in ["top", "left", "bottom", "right"]:
                    elem = OxmlElement(f"w:{edge}")
                    elem.set(qn("w:val"), "single")
                    elem.set(qn("w:sz"), "4")
                    elem.set(qn("w:color"), "D9D9D9")
                    borders.append(elem)
                tcpr.append(borders)
                margins = OxmlElement("w:tcMar")
                for edge, amount in [("top", "80"), ("bottom", "80"), ("left", "90"), ("right", "90")]:
                    elem = OxmlElement(f"w:{edge}")
                    elem.set(qn("w:w"), amount)
                    elem.set(qn("w:type"), "dxa")
                    margins.append(elem)
                tcpr.append(margins)
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.0
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        run.font.size = Pt(10)
                        run.font.bold = i == 0
        doc.add_paragraph().paragraph_format.space_after = Pt(0)
        return tab

    pct = lambda n: f"{n * 100:.2f}%"
    doc.add_paragraph("GST RAG Benchmark Findings", "Title")
    p("Experiment 01 | Source audit revision | 19 October 2026", small=True)
    p("Hybrid chunking is the provisional candidate for the next stage. It retrieved all labeled evidence within the top five results on the 18 answerable development questions. Parent expansion matched child-only coverage; neighboring-context expansion reduced coverage under the same token budget.")
    p("These are development retrieval findings from one chapter, rescored after an assistant audit of the complete PDF. Two ingestion defects remain open. The scores do not measure chatbot answer accuracy or establish correctness of current GST advice.")

    heading("Experiment scope")
    p(f"Input: one {manifest['page_count']}-page firm-book chapter on composite and mixed supply. Docling parsed the PDF once, and all methods used the same filtered source, Markdown table format, heading context, and Jina-v3 tokenizer.")
    p("The question set contains 30 source-audited questions: 20 development and 10 held out. Development has 18 answerable questions, one out-of-scope question, and one ambiguous question. The last two are excluded from recall averages. Held-out retrieval remains unscored. This is not an independent legal review.")

    heading("Chunking comparison")
    p("Fixed windows use a 512-token maximum with 64-token overlap. Hierarchical chunks follow Docling document elements. Hybrid chunks add token-aware splitting and merging, with a 512-token maximum.")
    rows = []
    for key, label in [("fixed", "Fixed size"), ("hierarchical", "Hierarchical"), ("hybrid", "Hybrid")]:
        s = main_run["variants"][key]["summary"]
        complete = sum(r["complete_at_10"] is True for r in main_run["variants"][key]["rows"])
        rows.append([label, pct(s["evidence_recall_at_5"]), pct(s["evidence_recall_at_10"]), f"{complete}/18", f"{s['reciprocal_rank_at_10']:.4f}"])
    table(["Method", "Recall at 5", "Recall at 10", "Complete at 10", "MRR at 10"], rows, [1.3, 1.3, 1.3, 1.85, 1.25])
    p("Recall averages the fraction of required evidence groups found per question. Complete evidence means every labeled group was found. MRR measures how early the first useful passage appears; it does not establish completeness.", small=True)

    doc.add_page_break()
    heading("Parent and child retrieval")
    p(f"The pipeline created {report['parents']} parents from explicitly numbered subsections. Children retain their parent IDs, source item references, and PDF pages. Parents include the subsection's subordinate examples and notes; cases are not forced into facts or ruling fields.")
    p("The follow-up experiment uses the same ranked hybrid children and a maximum of 2,048 context tokens. Only the assembly of the evidence changes.")
    p("Child only adds ranked passages. Neighbor expansion adds adjacent source elements within the same subsection. Parent expansion includes the full subsection when it fits, deduplicates repeated parents, and falls back to the matching child when a parent is too large.")
    context_rows = []
    for label, run in [("Child only", main_run), ("Neighboring elements", neighbors), ("Parent with child fallback", parents)]:
        s = run["variants"]["hybrid"]["summary"]
        context_rows.append([label, pct(s["context_evidence_recall"]), f"{s['context_tokens']:,.2f}"])
    table(["Context mode", "Evidence recall in context", "Mean context tokens"], context_rows, [3.25, 1.9, 1.85])
    p("Neighbor expansion crowded out required evidence for development question q05. Parent expansion did not improve coverage because the child-only context already contained all labeled evidence. Larger context is therefore not automatically more useful.")

    heading("What the implementation preserved")
    table(["Property", "Observed result"], [
        ["Chunks by method", "80 fixed; 244 hierarchical; 129 hybrid"],
        ["Source references covered", "All 285 canonical item references in every variant"],
        ["Missing PDF provenance", "Zero chunks"],
        ["Largest hierarchical element", "590 tokens; two elements exceed 512 tokens"],
        ["Largest fixed and hybrid chunks", "512 tokens including heading context"],
    ], [3.25, 3.75])
    p("Canonical-reference coverage does not include every visual element: the page-21 flowchart is missing from search. Two trailing-dot headings also need regrouping. PDF and printed pages were visually matched during audit, but that mapping is not implemented in ingestion.", small=True)

    doc.add_page_break()
    heading("Technical controls and verification")
    p("Jina's jina-embeddings-v3 model embeds passages with retrieval.passage and questions with retrieval.query. Normalized vectors are ranked locally using exact cosine similarity. Keeping ranking local removes approximate-index settings from this chunking experiment; Pinecone integration remains planned.")
    p("Evidence matching requires source-reference overlap and the labeled excerpt in the same retrieved record after whitespace and quote normalization. It is a conservative lexical proxy, not claim entailment. It does not bind the quote to a specific item span inside a merged record.")
    p("Docling detected headings at uniform depth in this PDF. Explicit numbering was used to reconstruct levels, and adjustments were saved separately from the original conversion. The same cleanup and heading correction were applied to every method.")
    p("Ten contract tests passed. They cover evidence matching, independent evidence groups, permission filtering, parent deduplication, large-parent fallback, neighbor boundaries, and fixed-window provenance. Artifact checks confirmed reference coverage and context budgets. A Plus-policy smoke run retrieved zero firm-book chunks; authentication is not yet implemented.")

    heading("Usage and reproducibility")
    p(f"The initial embedding run used {usage['tokens']:,} provider-reported tokens across {usage['requests']} requests, with {usage['seconds']:.3f} seconds of aggregate request time. Cached reruns reused vectors. This is not end-to-end answer latency or a monetary cost estimate.")
    p("Dependency versions are locked. The run records Docling 2.140.0, docling-core 2.102.0, Python 3.12.14, and configuration ID 749f74486094. Hosted model internals are not pinned, so future uncached API calls may differ.")

    heading("Decision and next steps")
    p("Keep hybrid and child-only context provisional. The assistant source audit is complete; repair subsection grouping and the missing diagram, add qualification-focused regression questions, then freeze the revised configuration before held-out scoring. Add an ICAI source and notification before evaluating updates across sources.")
    p("The corpus is small and some questions resemble source wording. No answer generator, independent legal-answer assessment, abstention policy, frontend, or admin-review application has been evaluated.")

    heading("Interview explanation")
    p('"I compared three chunking methods on one Docling parse using Jina embeddings. Hybrid performed best on the development evidence labels. A full assistant source audit then found stronger label requirements and two ingestion defects that those scores had not exposed. I keep those limitations explicit; held-out evaluation remains pending."')
    p("Source records: docs/SOURCE_AUDIT.md; manifest.json and chunk_report.json; source_audited_v2 development rescores. API usage comes from the initial benchmark_jina_dev_child_pro.json run. Original draft labels and results are preserved.", small=True)

    doc.add_page_break()
    heading("Complete source audit")
    p("The assistant read all 45 original PDF pages as text and page renders, including footnotes, tables, and the flowchart, and audited all 30 labels. Seventeen questions changed: sixteen now require more complete evidence, and one accepts equivalent repetitions. Required groups increased from 37 to 58; question wording and splits did not change.")
    p("Examples include the full conditions of composite supply, the ancillary relationship in principal supply, and the reason for the author's disagreement with a ruling. Retrieval of a phrase saying a ruling was wrong is not enough to support an explanation of why.")
    p("The revised scores reuse exactly the earlier saved development hits and contexts. No new embeddings or retrieval were run. Hierarchical recall at five rose from 83.33% to 87.50% because corrected labels can recognize partial evidence differently; this is a measurement change, not a retriever improvement. Hybrid still covers all development groups.")

    heading("Open defects and required verification")
    table(["Issue", "What must pass before calling it fixed"], [
        ["Trailing-dot headings", "Recognize 7.2.10. and 7.2.11. and verify rebuilt units belong to those parents, not 7.2.9."],
        ["Page 21 flowchart", "Verify a searchable transcription against every diagram branch, retain picture/page provenance, and test retrieval."],
        ["Later qualifications", "Add development questions requiring an original position and its relevant later appeal, note, or amendment."],
    ], [1.6, 5.4])
    p("The current readiness command intentionally fails on the heading and diagram defects. Existing contract tests pass, but they do not yet cover those repairs. Rebuild all methods against the same corrected source under a new configuration version, then compare both per-question misses and aggregate evidence coverage.")

    heading("Source ambiguities to preserve")
    p("Page 15 describes a dated renewable-device rate increase but later uses the old rate without clear temporal scope. Page 32 gives two different circular numbers. Pages 35-37 distinguish the author's hospital opinions from rulings and amendments. Preserve attribution and uncertainty rather than silently rewriting the source.")

    heading("Evaluation order")
    p("First reproduce each defect, then apply its repair and require the targeted regression to pass without breaking prior valid cases. Next freeze the revised source, labels, and settings, and score held-out retrieval without using its misses to retune. When generation is added, evaluate claim support, citation pages, opinion attribution, insufficient-evidence behavior, and manual review separately.")
    p("The full page-by-page ledger, per-question audit notes, reproduction commands, and unresolved checks are recorded in docs/SOURCE_AUDIT.md and local data/audit artifacts. Source SHA256 begins bddb8207104b9d46. The original PDF was not modified.", small=True)

    output = ROOT / "output/documents/GST_RAG_Benchmark_Findings_Source_Audit.docx"
    output.parent.mkdir(parents=True, exist_ok=True)
    # Remove decorative paragraph borders inherited from the bundled template.
    for style in doc.styles:
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.save(output)
    print(output)


if __name__ == "__main__":
    main()
